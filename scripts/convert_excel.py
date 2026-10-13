"""Konversi data mentah Excel (dataset Power BI) menjadi JSON untuk aplikasi web.

Cara pakai (dari folder project):
    pip install openpyxl
    python scripts/convert_excel.py data/raw/Dummy_Maintenance_Dashboard_Data.xlsx src/data/site-a.json

Yang dilakukan skrip ini:
1. Membaca sheet Data_WO, PM_Master, dan PM_Compliance.
2. Menyamakan nama kolom dengan format aplikasi (unitCode, type, status, ...).
3. Mengubah "Tanggal + Downtime_Jam" menjadi waktu mulai & selesai.
   WO pada unit dan tanggal yang sama dikerjakan berurutan mulai pukul 07:00,
   supaya downtime-nya tidak tumpang tindih.
"""

import json
import sys
from datetime import datetime, timedelta

from openpyxl import load_workbook

HEADER_ROW = 4  # baris judul kolom di setiap sheet (baris 1-3 berisi judul sheet)
STATUS_MAP = {"Completed": "Completed", "In Progress": "In Progress", "Open": "Open"}


def read_table(sheet):
    """Baca tabel mulai dari baris judul, berhenti di baris kosong pertama."""
    rows = sheet.iter_rows(min_row=HEADER_ROW, values_only=True)
    header = [h for h in next(rows) if h]
    table = []
    for row in rows:
        if row[0] is None or row[0] == "RINGKASAN":
            break
        table.append(dict(zip(header, row)))
    return table


def fmt(dt):
    return dt.strftime("%Y-%m-%dT%H:%M")


def main(xlsx_path, json_path):
    wb = load_workbook(xlsx_path, data_only=True)

    units = [
        {
            "code": r["Unit_ID"],
            "siteId": "A",
            "type": r["Tipe_Alat"],
            "brand": None,
            "model": None,
            "hmEnd": r["HM_Saat_Ini"],
            "pmInterval": r["Interval_PM_Jam"],
            "ewh": r["EWH_HM_per_Hari"],
        }
        for r in read_table(wb["PM_Master"])
    ]

    work_orders = []
    next_start = {}  # (unit, tanggal) -> waktu mulai WO berikutnya
    for r in read_table(wb["Data_WO"]):
        key = (r["Unit_ID"], r["Tanggal"])
        start = next_start.get(key, datetime.strptime(r["Tanggal"], "%Y-%m-%d") + timedelta(hours=7))
        end = start + timedelta(hours=float(r["Downtime_Jam"]))
        next_start[key] = end
        work_orders.append(
            {
                "id": r["No_WO"],
                "siteId": "A",
                "unitCode": r["Unit_ID"],
                "type": r["Jenis_Maintenance"],
                "status": STATUS_MAP[r["Status_WO"]],
                "start": fmt(start),
                "end": fmt(end),
                "component": None,
                "problem": None,
                "delayReason": None,
                "delayHours": 0,
                "cost": int(r["Biaya_Rp"]),
                "source": "excel",
            }
        )

    pm_logs = [
        {
            "unitCode": r["Unit_ID"],
            "pmType": r["Jenis_PM"],
            "hmPlan": r["HM_Jadwal"],
            "hmActual": r["HM_Aktual"],
        }
        for r in read_table(wb["PM_Compliance"])
    ]

    data = {
        "source": xlsx_path.split("/")[-1],
        "cutoff": "2025-12-31",
        "units": units,
        "workOrders": work_orders,
        "pmLogs": pm_logs,
    }
    # Satu record per baris: ringkas tapi tetap mudah dibaca/di-diff
    def records(items):
        return ",\n".join("    " + json.dumps(item, ensure_ascii=False) for item in items)

    with open(json_path, "w", encoding="utf-8") as f:
        f.write("{\n")
        f.write(f'  "source": {json.dumps(data["source"])},\n')
        f.write(f'  "cutoff": {json.dumps(data["cutoff"])},\n')
        for i, key in enumerate(["units", "workOrders", "pmLogs"]):
            comma = "," if i < 2 else ""
            f.write(f'  "{key}": [\n{records(data[key])}\n  ]{comma}\n')
        f.write("}\n")
    print(f"{len(units)} unit, {len(work_orders)} work order, {len(pm_logs)} log PM -> {json_path}")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
